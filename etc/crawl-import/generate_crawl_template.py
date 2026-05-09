# One-off script: generates KHHub crawl/import Excel template.
# Run from repo root: python3 etc/crawl-import/generate_crawl_template.py

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def header_row(ws, headers, row=1):
    fill = PatternFill("solid", fgColor="3B82F6")
    font = Font(color="FFFFFF", bold=True, size=11)
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = fill
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")


def sample_row(ws, values, row=2):
    for col, v in enumerate(values, 1):
        ws.cell(row=row, column=col, value=v)


def autosize_columns(ws, max_width=55):
    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)), max_width))
        ws.column_dimensions[col_letter].width = max(12, min(max_len + 2, max_width))


def main():
    root = Path(__file__).resolve().parent
    out = root / "KHHub_Crawl_Import_Template.xlsx"

    wb = Workbook()
    # --- Instructions sheet
    ws0 = wb.active
    ws0.title = "HuongDan"
    instructions = [
        ("Mục đích", "File mẫu để crawl dữ liệu Articles / Places / Jobs rồi import vào KH Hub MasterData."),
        ("Sheet", "Mỗi loại entity một sheet: Articles_Crawl, Places_Crawl, Jobs_Crawl."),
        ("Category", "Điền CategorySlug hoặc CategoryName trùng với danh mục đã tạo trong admin (import sẽ map sang Guid)."),
        ("TagsSuggested", "Gợi ý tag cho bản ghi crawl: phân cách bằng dấu phẩy hoặc dấu chấm phẩy. Import có thể tách, chuẩn hoá slug, tạo tag nếu chưa có."),
        ("Enum (C#)", "Giá trị enum viết đúng tên như trong code (ví dụ Published, FullTime)."),
        ("Province / Ward", "Nên dùng tên hoặc mã thống nhất với bảng Provinces/Wards trong DB để resolver khi import."),
        ("Slug", "URL-friendly, không dấu, chữ thường, gạch ngang."),
        ("Content / Description", "Có thể là HTML hoặc plain text tùy pipeline import."),
        ("ExternalId / CrawlSourceUrl", "Tuỳ chọn: để trùng lặp và truy vết nguồn crawl."),
        ("", ""),
        ("ArticleType", "News, Blog, Guide, Review, Promotion, Event"),
        ("ArticleStatus", "Draft, Pending, Published, Archived"),
        ("PlaceStatus", "Draft, Published, Hidden, Archived"),
        ("PriceRange", "Free, Cheap, Medium, High, Luxury"),
        ("JobStatus", "Draft, Published, Closed, Expired, Archived"),
        ("EmploymentType", "FullTime, PartTime, Internship, Freelance, Contract"),
        ("WorkMode", "Onsite, Hybrid, Remote"),
        ("ExperienceLevel", "Fresher, Junior, Middle, Senior, Lead"),
    ]
    ws0["A1"] = "Key"
    ws0["B1"] = "Value / ghi chú"
    ws0["A1"].font = Font(bold=True)
    ws0["B1"].font = Font(bold=True)
    for i, (k, v) in enumerate(instructions, start=2):
        ws0.cell(row=i, column=1, value=k)
        ws0.cell(row=i, column=2, value=v)
    ws0.column_dimensions["A"].width = 22
    ws0.column_dimensions["B"].width = 92

    # --- Articles
    ws1 = wb.create_sheet("Articles_Crawl")
    art_headers = [
        "ExternalId",
        "CrawlSourceUrl",
        "CategorySlug",
        "CategoryName",
        "TagsSuggested",
        "Title",
        "Slug",
        "Summary",
        "Content",
        "ThumbnailUrl",
        "CoverImageUrl",
        "Type",
        "AuthorName",
        "Source",
        "SourceUrl",
        "Status",
        "PublishedAt",
        "IsFeatured",
        "IsHot",
        "IsTrending",
        "ViewCount",
        "LikeCount",
        "ShareCount",
        "CommentCount",
        "ReadingTime",
        "SeoTitle",
        "SeoDescription",
        "SeoKeywords",
    ]
    header_row(ws1, art_headers)
    sample_row(
        ws1,
        [
            "ext-article-001",
            "https://example.com/news/sample",
            "tin-tuc",
            "",
            "Hà Nội, du lịch, ẩm thực",
            "Tiêu đề bài viết mẫu",
            "tieu-de-bai-viet-mau",
            "Tóm tắt ngắn 1–2 câu.",
            "<p>Nội dung HTML hoặc plain text.</p>",
            "",
            "",
            "News",
            "Tác giả crawl",
            "Nguồn gốc",
            "https://source.example.com/article/1",
            "Draft",
            "2026-05-09T10:00:00",
            "FALSE",
            "FALSE",
            "FALSE",
            "0",
            "0",
            "0",
            "0",
            "0",
            "Tiêu đề bài viết mẫu",
            "Mô tả SEO",
            "từ khoá, seo",
        ],
    )
    autosize_columns(ws1)

    # --- Places
    ws2 = wb.create_sheet("Places_Crawl")
    place_headers = [
        "ExternalId",
        "CrawlSourceUrl",
        "CategorySlug",
        "CategoryName",
        "TagsSuggested",
        "Name",
        "Slug",
        "ShortDescription",
        "Description",
        "ThumbnailUrl",
        "CoverImageUrl",
        "Address",
        "Latitude",
        "Longitude",
        "PhoneNumber",
        "Email",
        "Website",
        "OpeningHours",
        "PriceRange",
        "GoogleMapUrl",
        "Status",
        "ProvinceNameOrCode",
        "WardNameOrCode",
        "ViewCount",
        "FavoriteCount",
        "ReviewCount",
        "RatingAveraged",
        "RatingTotal",
        "IsFeatured",
        "IsHot",
        "IsVerified",
        "SeoTitle",
        "SeoDescription",
        "SeoKeywords",
    ]
    header_row(ws2, place_headers)
    sample_row(
        ws2,
        [
            "ext-place-001",
            "https://maps.example.com/p/123",
            "nha-hang",
            "",
            "ăn uống, trung tâm, view đẹp",
            "Quán cà phê mẫu",
            "quan-ca-phe-mau",
            "Mô tả ngắn.",
            "Mô tả chi tiết HTML hoặc text.",
            "",
            "",
            "123 Đường mẫu, Quận mẫu",
            "21.0285",
            "105.8542",
            "",
            "",
            "",
            "08:00-22:00",
            "Medium",
            "",
            "Draft",
            "Hà Nội",
            "Phường Tràng Tiền",
            "0",
            "0",
            "0",
            "0",
            "0",
            "FALSE",
            "FALSE",
            "FALSE",
            "Quán cà phê mẫu",
            "SEO mô tả địa điểm",
            "cafe, hà nội",
        ],
    )
    autosize_columns(ws2)

    # --- Jobs
    ws3 = wb.create_sheet("Jobs_Crawl")
    job_headers = [
        "ExternalId",
        "CrawlSourceUrl",
        "CategorySlug",
        "CategoryName",
        "TagsSuggested",
        "Title",
        "Slug",
        "Summary",
        "Description",
        "Requirements",
        "Benefits",
        "ThumbnailUrl",
        "CoverImageUrl",
        "EmploymentType",
        "WorkMode",
        "ExperienceLevel",
        "SalaryMin",
        "SalaryMax",
        "SalaryText",
        "SalaryCurrency",
        "Location",
        "ContactEmail",
        "ContactPhone",
        "ApplicationUrl",
        "PublishedAt",
        "Status",
        "ProvinceNameOrCode",
        "WardNameOrCode",
        "ViewCount",
        "ApplicationCount",
        "FavoriteCount",
        "ShareCount",
        "IsFeatured",
        "IsUrgent",
        "IsHot",
        "SeoTitle",
        "SeoDescription",
        "SeoKeywords",
    ]
    header_row(ws3, job_headers)
    sample_row(
        ws3,
        [
            "ext-job-001",
            "https://jobs.example.com/j/999",
            "it-phan-mem",
            "",
            "Laravel, PHP, remote",
            "Senior Backend Developer",
            "senior-backend-developer",
            "Tóm tắt tin tuyển dụng.",
            "<p>Mô tả công việc...</p>",
            "Yêu cầu 1...\nYêu cầu 2...",
            "Lương cạnh tranh, BHXH...",
            "",
            "",
            "FullTime",
            "Hybrid",
            "Senior",
            "30000000",
            "50000000",
            "30–50 triệu",
            "VND",
            "Hà Nội",
            "hr@company.example",
            "",
            "https://company.example/apply",
            "2026-05-09T10:00:00",
            "Draft",
            "Hà Nội",
            "Phường Tràng Tiền",
            "0",
            "0",
            "0",
            "0",
            "FALSE",
            "FALSE",
            "FALSE",
            "Senior Backend Developer",
            "Tuyển Senior Backend",
            "backend, php, laravel",
        ],
    )
    autosize_columns(ws3)

    wb.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
